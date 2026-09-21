from flask import Flask, render_template, request, redirect, url_for, session, flash, jsonify, send_file
from werkzeug.security import generate_password_hash, check_password_hash
from pathlib import Path
from functools import wraps
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import sqlite3, os, io, json

BASE = Path(__file__).resolve().parent
DB_DIR = BASE / "instance"
UPLOAD_DIR = BASE / "static" / "uploads"
DB_DIR.mkdir(exist_ok=True)
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
DB = DB_DIR / "scgpc_lab.db"

app = Flask(__name__)
app.secret_key = os.getenv("SECRET_KEY", "change-this-secret-in-render")

PROJECT = "Effect of Sugarcane Bagasse Ash (SCBA) on Fire Resistance of Fly Ash–GGBS Based SCGPC"
MIXES = [
    {"id":"M0","scba":0,"molarity":"4M"},{"id":"M1","scba":5,"molarity":"4M"},
    {"id":"M2","scba":10,"molarity":"4M"},{"id":"M3","scba":15,"molarity":"4M"},
    {"id":"M4","scba":20,"molarity":"4M"},{"id":"M5","scba":25,"molarity":"4M"},
    {"id":"M6","scba":0,"molarity":"6M"},{"id":"M7","scba":5,"molarity":"6M"},
    {"id":"M8","scba":10,"molarity":"6M"},{"id":"M9","scba":15,"molarity":"6M"},
    {"id":"M10","scba":20,"molarity":"6M"},{"id":"M11","scba":25,"molarity":"6M"}
]
TEMPS = [200,400,600,800]
ROLES = ["user","editor","superadmin"]

def conn():
    c = sqlite3.connect(DB)
    c.row_factory = sqlite3.Row
    return c

def init():
    c=conn()
    c.executescript("""
    CREATE TABLE IF NOT EXISTS users(
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      username TEXT UNIQUE NOT NULL,
      password_hash TEXT NOT NULL,
      display_name TEXT NOT NULL,
      role TEXT NOT NULL
    );
    CREATE TABLE IF NOT EXISTS specimens(
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      mix_id TEXT NOT NULL,
      temp INTEGER NOT NULL,
      replicate INTEGER NOT NULL,
      test_date TEXT,
      pre_mass REAL,
      post_mass REAL,
      failure_load REAL,
      baseline_strength REAL,
      colour_change TEXT,
      cracks TEXT,
      spalling TEXT,
      heating_rate TEXT,
      exposure_duration TEXT,
      furnace TEXT,
      remarks TEXT,
      image_path TEXT,
      before_image TEXT,
      after_image TEXT,
      status TEXT DEFAULT 'Pending',
      updated_by TEXT,
      updated_at TEXT DEFAULT CURRENT_TIMESTAMP,
      UNIQUE(mix_id,temp,replicate)
    );
    CREATE TABLE IF NOT EXISTS notes(
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      title TEXT NOT NULL,
      body TEXT NOT NULL,
      author TEXT,
      created_at TEXT DEFAULT CURRENT_TIMESTAMP
    );
    """)
    defaults=[
      ("abhishek_admin",os.getenv("SUPERADMIN_PASSWORD","SCGPC@Admin2026"),"Abhishek","superadmin"),
      ("projectlead",os.getenv("EDITOR_PASSWORD","SCGPC@Lead2026"),"Project Lead / Faculty","editor"),
      ("common_user",os.getenv("USER_PASSWORD","SCGPC@User2026"),"Common User","user")
    ]
    for u,p,n,r in defaults:
        if not c.execute("SELECT id FROM users WHERE username=?",(u,)).fetchone():
            c.execute("INSERT INTO users(username,password_hash,display_name,role) VALUES(?,?,?,?)",
                      (u,generate_password_hash(p),n,r))
    for m in MIXES:
        for t in TEMPS:
            for rep in (1,2,3):
                c.execute("""INSERT OR IGNORE INTO specimens(mix_id,temp,replicate)
                             VALUES(?,?,?)""",(m["id"],t,rep))
    c.commit(); c.close()

init()

def user():
    return session.get("user")

def auth(fn):
    @wraps(fn)
    def w(*a,**kw):
        if not user(): return redirect(url_for("login",next=request.path))
        return fn(*a,**kw)
    return w

def allow(*roles):
    def dec(fn):
        @wraps(fn)
        def w(*a,**kw):
            if not user(): return redirect(url_for("login"))
            if user()["role"] not in roles:
                flash("Access restricted for this account.","error")
                return redirect(url_for("dashboard"))
            return fn(*a,**kw)
        return w
    return dec

@app.context_processor
def common():
    return {"me":user(),"mixes":MIXES,"temps":TEMPS,"project":PROJECT}

@app.route("/")
def index(): return redirect(url_for("dashboard") if user() else url_for("login"))

@app.route("/login",methods=["GET","POST"])
def login():
    if request.method=="POST":
        u=request.form.get("username","").strip()
        p=request.form.get("password","")
        c=conn(); row=c.execute("SELECT * FROM users WHERE username=?",(u,)).fetchone(); c.close()
        if row and check_password_hash(row["password_hash"],p):
            session["user"]={"id":row["id"],"username":row["username"],"display_name":row["display_name"],"role":row["role"]}
            return redirect(request.args.get("next") or url_for("dashboard"))
        flash("Invalid username or password.","error")
    return render_template("login.html")

@app.route("/logout")
def logout():
    session.clear(); return redirect(url_for("login"))

def calc(r):
    mass=None; residual=None; pct=None
    if r["pre_mass"] is not None and r["post_mass"] is not None and r["pre_mass"] != 0:
        mass=(r["pre_mass"]-r["post_mass"])/r["pre_mass"]*100
    if r["failure_load"] is not None: residual=r["failure_load"]/10
    if residual is not None and r["baseline_strength"]:
        pct=residual/r["baseline_strength"]*100
    return mass,residual,pct

def complete(r):
    return all(r[k] is not None for k in ("test_date","pre_mass","post_mass","failure_load","baseline_strength"))

@app.route("/dashboard")
@auth
def dashboard():
    c=conn()
    rows=c.execute("SELECT * FROM specimens").fetchall(); c.close()
    completed=sum(complete(r) for r in rows)
    total=len(rows); remaining=total-completed
    bytemp={t:sum(complete(r) for r in rows if r["temp"]==t) for t in TEMPS}
    bymix={m["id"]:sum(complete(r) for r in rows if r["mix_id"]==m["id"]) for m in MIXES}
    return render_template("dashboard.html",completed=completed,total=total,remaining=remaining,
                           bytemp=bytemp,bymix=bymix,pct=(completed/total*100 if total else 0))

@app.route("/project")
@auth
def project_page():
    return render_template("project.html")

@app.route("/data-entry")
@auth
def data_entry():
    mix=request.args.get("mix",""); temp=request.args.get("temp","")
    c=conn(); rows=c.execute("SELECT * FROM specimens ORDER BY mix_id,temp,replicate").fetchall(); c.close()
    return render_template("data_entry.html",rows=rows,mix_filter=mix,temp_filter=temp)

@app.route("/save-specimen",methods=["POST"])
@auth
def save_specimen():
    sid=request.form.get("id")
    c=conn(); old=c.execute("SELECT * FROM specimens WHERE id=?",(sid,)).fetchone()
    if not old: c.close(); flash("Specimen not found.","error"); return redirect(url_for("data_entry"))
    if user()["role"]=="user" and old["status"]=="Completed":
        c.close(); flash("This completed record is locked. Ask an Editor or Super Admin to edit it.","error")
        return redirect(url_for("data_entry"))
    def n(k):
        x=request.form.get(k,"").strip()
        return float(x) if x else None
    image=old["image_path"]
    f=request.files.get("after_image")
    if f and f.filename:
        name=f"{old['mix_id']}_{old['temp']}C_R{old['replicate']}_{os.path.basename(f.filename).replace(' ','_')}"
        f.save(UPLOAD_DIR/name); image=f"uploads/{name}"
    vals=(request.form.get("test_date") or None,n("pre_mass"),n("post_mass"),n("failure_load"),
          n("baseline_strength"),request.form.get("colour_change"),request.form.get("cracks"),
          request.form.get("spalling"),request.form.get("heating_rate"),request.form.get("exposure_duration"),
          request.form.get("furnace"),request.form.get("remarks"),image,request.form.get("status","Pending"),
          user()["username"],sid)
    c.execute("""UPDATE specimens SET test_date=?,pre_mass=?,post_mass=?,failure_load=?,baseline_strength=?,
      colour_change=?,cracks=?,spalling=?,heating_rate=?,exposure_duration=?,furnace=?,remarks=?,
      image_path=?,status=?,updated_by=?,updated_at=CURRENT_TIMESTAMP WHERE id=?""",vals)
    c.commit(); c.close()
    flash("Specimen saved successfully.","success")
    return redirect(url_for("data_entry",mix=old["mix_id"],temp=old["temp"]))

@app.route("/graphs")
@auth
def graphs(): return render_template("graphs.html")

@app.route("/api/data")
@auth
def api_data():
    c=conn(); rows=c.execute("SELECT * FROM specimens ORDER BY mix_id,temp,replicate").fetchall(); c.close()
    out=[]
    for r in rows:
        m,s,p=calc(r)
        out.append({"mix":r["mix_id"],"temp":r["temp"],"rep":r["replicate"],"mass_loss":m,
                    "residual":s,"residual_pct":p,"scba":next(x["scba"] for x in MIXES if x["id"]==r["mix_id"]),
                    "molarity":next(x["molarity"] for x in MIXES if x["id"]==r["mix_id"])})
    return jsonify(out)

@app.route("/results")
@auth
def results():
    c=conn(); rows=c.execute("SELECT * FROM specimens ORDER BY mix_id,temp,replicate").fetchall(); c.close()
    return render_template("results.html",rows=rows,calc=calc,complete=complete)

@app.route("/notes",methods=["GET","POST"])
@auth
def notes():
    c=conn()
    if request.method=="POST":
        c.execute("INSERT INTO notes(title,body,author) VALUES(?,?,?)",
                  (request.form.get("title"),request.form.get("body"),user()["display_name"]))
        c.commit(); flash("Note added.","success")
    rows=c.execute("SELECT * FROM notes ORDER BY id DESC").fetchall(); c.close()
    return render_template("notes.html",notes=rows)

@app.route("/export")
@auth
def export():
    c=conn(); rows=c.execute("SELECT * FROM specimens ORDER BY mix_id,temp,replicate").fetchall(); c.close()
    wb=Workbook(); ws=wb.active; ws.title="Fire Data"
    heads=["Specimen","Mix","SCBA %","NaOH","Temperature (°C)","Replicate","Test Date",
           "Pre-fire Mass (g)","Post-fire Mass (g)","Mass Loss (%)","Failure Load (kN)",
           "Residual Strength (MPa)","28D Baseline (MPa)","Residual Strength (%)",
           "Colour Change","Cracks","Spalling","Heating Rate","Exposure Duration","Furnace",
           "Remarks","Status","Updated By"]
    ws.append(heads)
    for c0 in ws[1]:
        c0.font=Font(bold=True,color="FFFFFF"); c0.fill=PatternFill("solid",fgColor="103D5B")
    mm={x["id"]:x for x in MIXES}
    for r in rows:
        m,s,p=calc(r); x=mm[r["mix_id"]]
        ws.append([f"{r['mix_id']}-{r['temp']}C-R{r['replicate']}",r["mix_id"],x["scba"],x["molarity"],r["temp"],r["replicate"],
                   r["test_date"],r["pre_mass"],r["post_mass"],m,r["failure_load"],s,r["baseline_strength"],p,
                   r["colour_change"],r["cracks"],r["spalling"],r["heating_rate"],r["exposure_duration"],r["furnace"],
                   r["remarks"],"Completed" if complete(r) else "Pending",r["updated_by"]])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width=min(max(len(str(col[0].value or ""))+3,12),28)
    ws.freeze_panes="A2"
    bio=io.BytesIO(); wb.save(bio); bio.seek(0)
    return send_file(bio,as_attachment=True,download_name="SCGPC_Fire_Resistance_Results.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.route("/admin",methods=["GET","POST"])
@allow("superadmin")
def admin():
    c=conn()
    if request.method=="POST":
        u=request.form.get("username","").strip(); p=request.form.get("password","")
        n=request.form.get("display_name","").strip(); r=request.form.get("role","user")
        if u and p and n and r in ROLES:
            try:
                c.execute("INSERT INTO users(username,password_hash,display_name,role) VALUES(?,?,?,?)",
                          (u,generate_password_hash(p),n,r)); c.commit(); flash("User created.","success")
            except sqlite3.IntegrityError: flash("Username already exists.","error")
    users=c.execute("SELECT id,username,display_name,role FROM users ORDER BY id").fetchall(); c.close()
    return render_template("admin.html",users=users)

@app.route("/admin/delete-user/<int:uid>",methods=["POST"])
@allow("superadmin")
def delete_user(uid):
    if uid==user()["id"]: flash("You cannot delete your own account.","error"); return redirect(url_for("admin"))
    c=conn(); c.execute("DELETE FROM users WHERE id=?",(uid,)); c.commit(); c.close()
    flash("User deleted.","success"); return redirect(url_for("admin"))

if __name__=="__main__":
    app.run(host="0.0.0.0",port=int(os.getenv("PORT",5000)),debug=True)
